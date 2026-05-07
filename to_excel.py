import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

src = r"C:\Users\Rares\Desktop\testATCA\receipts_classified.csv"
dst = r"C:\Users\Rares\Desktop\testATCA\receipts_classified.xlsx"

wb = Workbook()
ws = wb.active
ws.title = "Bonuri fiscale"

with open(src, encoding="utf-8") as f:
    reader = csv.reader(f, delimiter=";")
    rows = list(reader)

for r_idx, row in enumerate(rows, start=1):
    for c_idx, val in enumerate(row, start=1):
        cell = ws.cell(row=r_idx, column=c_idx, value=val)
        if r_idx == 1:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="305496")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

# Convert numeric-looking columns to numbers
numeric_cols = {"Cantitate", "Pret unitar (net)", "Valoare (net)", "TVA", "Total"}
header = rows[0]
num_indices = [i + 1 for i, h in enumerate(header) if h in numeric_cols]

for r_idx in range(2, len(rows) + 1):
    for c_idx in num_indices:
        cell = ws.cell(row=r_idx, column=c_idx)
        v = cell.value
        if v in (None, "", "-"):
            continue
        try:
            cell.value = float(v)
            cell.number_format = "#,##0.00"
        except ValueError:
            pass

# Column widths
widths = {
    "Pagina": 7, "Furnizor": 28, "CIF Furnizor": 13, "Client": 38,
    "Data": 12, "Nr. document": 22, "Cod produs": 14, "Nume produs": 38,
    "Cantitate": 10, "UM": 6, "Pret unitar (net)": 14, "Valoare (net)": 13,
    "TVA": 10, "Total": 11, "Cont contabil": 10, "Observatii": 50,
}
for i, h in enumerate(header, start=1):
    ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 15)

ws.row_dimensions[1].height = 32
ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions

# Wrap text in Observatii column
obs_idx = header.index("Observatii") + 1
for r in range(2, len(rows) + 1):
    ws.cell(row=r, column=obs_idx).alignment = Alignment(wrap_text=True, vertical="top")

wb.save(dst)
print("Saved:", dst)
